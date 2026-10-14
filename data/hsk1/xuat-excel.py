# -*- coding: utf-8 -*-
"""
Đổi một file bài-XX.json thành file Excel nhập liệu cho khu quản trị.

    python xuat-excel.py bai-06.json excel/bai-06.xlsx

Một file Excel chứa đủ 8 sheet cho cả 6 phần soạn bài, nên nhập xong một phần thì chọn lại
đúng file đó cho phần sau. Tên sheet và tên cột phải khớp với web/src/lib/admin/import/*.ts —
bộ nhập tìm sheet theo tên trước, nên sai tên là nó lấy nhầm sheet đầu file.
"""

import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

# (tên sheet, khoá trong JSON, tiêu đề các cột)
SHEETS = [
    ('Từ vựng', 'tu_vung',
     ['Chữ Hán', 'Pinyin', 'Hán Việt', 'Nghĩa tiếng Việt', 'Loại từ', 'Ví dụ', 'Ví dụ pinyin', 'Ví dụ nghĩa']),
    ('Bài tập', 'bai_tap',
     ['Dạng câu', 'Nội dung', 'Phụ', 'A', 'B', 'C', 'D', 'Đáp án', 'Ô chữ', 'Câu mẫu', 'Giải thích', 'Điểm']),
    ('Kiểm tra', 'kiem_tra',
     ['Dạng câu', 'Nội dung', 'Phụ', 'A', 'B', 'C', 'D', 'Đáp án', 'Ô chữ', 'Câu mẫu', 'Giải thích', 'Điểm']),
    ('Luyện viết', 'luyen_viet',
     ['Chữ Hán', 'Pinyin', 'Hán Việt', 'Nghĩa', 'Cấu tạo', 'Hiện']),
    ('Lời thoại', 'loi_thoai',
     ['Bài khoá', 'Nhân vật', 'Vai', 'Chữ đại diện', 'Tiếng Trung', 'Pinyin', 'Nghĩa tiếng Việt', 'Từ nhấn']),
    ('Bài khoá', 'bai_khoa',
     ['Bài khoá', 'Tên bài khoá', 'Tình huống']),
    ('Ngữ pháp', 'ngu_phap',
     ['Tên cấu trúc', 'Mô tả ngắn', 'Giải thích', 'Lưu ý', 'Khối công thức', 'Câu luyện tập', 'Nghĩa luyện tập']),
    ('Ví dụ', 'vi_du_ngu_phap',
     ['Tên cấu trúc', 'Loại', 'Công thức', 'Tiếng Trung', 'Pinyin', 'Nghĩa', 'Từ nhấn']),
]

DAU = PatternFill('solid', fgColor='FFE3E8EF')
RONG = {'Giải thích': 46, 'Nội dung': 30, 'Tiếng Trung': 26, 'Pinyin': 26, 'Nghĩa tiếng Việt': 30,
        'Ví dụ': 26, 'Ví dụ pinyin': 26, 'Ví dụ nghĩa': 30, 'Ô chữ': 34, 'Khối công thức': 30,
        'Câu hỏi': 26, 'Tên cấu trúc': 22, 'Lưu ý': 34, 'Nghĩa': 22, 'Công thức': 26}


def main() -> None:
    nguon = Path(sys.argv[1])
    dich = Path(sys.argv[2])
    data = json.loads(nguon.read_text(encoding='utf-8'))

    wb = Workbook()
    wb.remove(wb.active)

    for ten, khoa, cot in SHEETS:
        ws = wb.create_sheet(ten)
        ws.append(cot)
        for o in ws[1]:
            o.font = Font(bold=True)
            o.fill = DAU
        for dong in data.get(khoa, []):
            ws.append(list(dong))
        for i, tieu_de in enumerate(cot, start=1):
            ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = RONG.get(tieu_de, 16)
        ws.freeze_panes = 'A2'
        for hang in ws.iter_rows(min_row=2):
            for o in hang:
                o.alignment = Alignment(vertical='top', wrap_text=True)

    dich.parent.mkdir(parents=True, exist_ok=True)
    wb.save(dich)

    print(f'Đã ghi {dich}')
    for ten, khoa, _ in SHEETS:
        print(f'  {ten:12} {len(data.get(khoa, []))} dòng')


if __name__ == '__main__':
    main()
